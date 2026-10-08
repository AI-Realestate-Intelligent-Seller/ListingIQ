"""CSV and XLSX ingestion for the lead pool.

Brokers buy lists from many vendors, so the column names are never the same
twice. Rather than demand one exact header, we recognise a set of aliases per
field and additionally treat any column named after a signal as a flag column
("fsbo,high_equity" with 1/yes/x values).
"""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Iterable

from .catalog import SIGNALS, normalize_signal, parse_signals

# Field -> header aliases, all compared lower-cased with separators stripped.
_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    'owner_name': ('owner', 'ownername', 'name', 'fullname', 'ownerfullname', 'contactname', 'seller'),
    'first_name': ('firstname', 'ownerfirstname', 'first'),
    'last_name': ('lastname', 'ownerlastname', 'last', 'surname'),
    'phone': ('phone', 'phonenumber', 'phone1', 'primaryphone', 'primaryphonenumber',
              'mobile', 'cell', 'cellphone', 'contact', 'contactnumber', 'telephone'),
    # Numbered contact fields are common in CRM exports. The first usable
    # number remains the lead's primary phone; the others travel in details.
    'phone_2': ('phone2', 'secondaryphone', 'secondaryphonenumber'),
    'phone_3': ('phone3', 'tertiaryphone', 'tertiaryphonenumber'),
    'property_address': ('propertyaddress', 'address', 'property', 'street', 'streetaddress', 'siteaddress'),
    'area': ('area', 'city', 'neighborhood', 'neighbourhood', 'market', 'town', 'submarket'),
    'state': ('state', 'property_state', 'propertystate'),
    'postal_code': ('zip', 'zipcode', 'postalcode', 'propertyzip', 'propertyzipcode'),
    'email_1': ('email', 'emailaddress', 'email1', 'primaryemail'),
    'email_2': ('email2', 'secondaryemail'),
    'ok_to_text': ('oktotext', 'permissiontotext', 'smspermission', 'smsconsent', 'textconsent'),
    'signals': ('signals', 'signal', 'tags', 'tag', 'leadtype', 'lead_type', 'category',
                'motivation', 'leadsource', 'source', 'listingstatus', 'status'),
    # Why this owner is being contacted, in the vendor's own words.
    'outreach_reason': ('outreachreason', 'reason', 'outreachnotes', 'note', 'notes'),
    # A JSON blob of property attributes (beds, baths, price, …).
    'details': ('leaddetails', 'details', 'propertydetails', 'propertyinfo', 'attributes'),
}

# Attributes inside a details blob that imply a signal. Only facts the record
# states outright — nothing is inferred beyond what the data says.
_DETAIL_SIGNALS: tuple[tuple[str, str, str], ...] = (
    ('listed_for_sale_by_owner', 'yes', 'fsbo'),
    ('owner_occ', 'no', 'absentee_owner'),
    ('vacant', 'yes', 'vacant'),
    ('pre_foreclosure', 'yes', 'pre_foreclosure'),
)

_TRUTHY = {'1', 'y', 'yes', 'true', 't', 'x', 'active'}

# Lofty puts useful CRM facts in separate columns rather than a JSON details
# cell. Preserve those facts in the existing details JSON; this intentionally
# requires no database/schema change.
_LOFTY_DETAIL_COLUMNS: dict[str, str] = {
    'Lofty Lead ID': 'lofty_lead_id',
    'Other Properties': 'other_properties',
    'Est. Home Value': 'estimated_value',
    'Purchase Date': 'purchase_date',
    'Lead Type': 'lead_type',
    'Selling Timeframe': 'selling_timeframe',
    'Has Listing Agent': 'has_listing_agent',
    'Opportunity': 'opportunity',
    'Lead Source': 'lead_source',
    'Lofty Stage': 'lofty_stage',
    'Assigned Agent': 'assigned_agent',
    'Ownership': 'ownership',
    'Created': 'crm_created_at',
    'Last Touch': 'last_touch',
    'Last Site Visit': 'last_site_visit',
    'Outreach Attempts (since 2024)': 'outreach_attempts_since_2024',
    'Texts Sent': 'texts_sent',
    'Calls Made': 'calls_made',
    'Emails Sent': 'emails_sent',
    'Email Opens': 'email_opens',
    'First Outreach': 'first_outreach',
    'Last Outreach': 'last_outreach',
    'Lead Score': 'crm_lead_score',
    'Language': 'language',
    'Level of Interest': 'level_of_interest',
    'Fello Link': 'fello_link',
    'Household Contacts': 'household_contacts',
    'Tags': 'crm_tags',
}


def _key(header: str) -> str:
    return re.sub(r'[^a-z0-9]+', '', str(header or '').lower())


# Fields a broker may map by hand when the guess is wrong.
MAPPABLE_FIELDS = (
    'owner_name', 'first_name', 'last_name',
    'phone', 'phone_2', 'phone_3', 'email_1', 'email_2',
    'property_address', 'area', 'state', 'postal_code',
    'signals', 'outreach_reason', 'details',
)


def _map_headers(fieldnames: list[str],
                 overrides: dict[str, str] | None = None) -> tuple[dict[str, str], dict[str, str]]:
    """Return (field -> column, signal key -> column) for one header row.

    An override wins over the guess, and an override of '' means "no column" —
    a broker must be able to say a detected column is the wrong one.
    """
    fields: dict[str, str] = {}
    flags: dict[str, str] = {}
    chosen: set[str] = set()

    for field, column in (overrides or {}).items():
        if field not in MAPPABLE_FIELDS:
            continue
        if column == '':
            # Explicitly unmapped: record nothing, and block the auto-guess.
            chosen.add(field)
            continue
        if column in (fieldnames or []):
            fields[field] = column
            chosen.add(field)

    for column in fieldnames or []:
        if column in fields.values():
            continue
        key = _key(column)
        if not key:
            continue
        for field, aliases in _FIELD_ALIASES.items():
            if field in fields or field in chosen:
                continue
            if key in {_key(alias) for alias in aliases}:
                fields[field] = column
                break
        else:
            signal = normalize_signal(column)
            if signal and signal not in flags:
                flags[signal] = column
    return fields, flags


def normalize_phone(value: str) -> str | None:
    """Best-effort E.164. North American numbers are assumed when unprefixed."""
    raw = str(value or '').strip()
    if not raw:
        return None
    plus = raw.startswith('+')
    digits = re.sub(r'\D', '', raw)
    if not digits:
        return None
    if plus:
        return f'+{digits}' if 7 <= len(digits) <= 15 else None
    if len(digits) == 10:
        return f'+1{digits}'
    if len(digits) == 11 and digits.startswith('1'):
        return f'+{digits}'
    return None


def parse_details(value: object) -> dict:
    """A property-details cell as a flat dict of strings.

    Vendors ship these as a JSON object in one column. Anything that is not an
    object is ignored rather than guessed at.
    """
    if isinstance(value, dict):
        parsed = value
    else:
        text = str(value or '').strip()
        if not text.startswith('{'):
            return {}
        try:
            parsed = json.loads(text)
        except ValueError:
            return {}
    if not isinstance(parsed, dict):
        return {}
    return {str(key): str(item) for key, item in parsed.items()
            if item is not None and str(item).strip() != ''}


def signals_from_details(details: dict) -> list[str]:
    """Signals the details blob states outright."""
    found = []
    for key, wanted, signal in _DETAIL_SIGNALS:
        if str(details.get(key, '')).strip().lower() == wanted and signal not in found:
            found.append(signal)
    return found


def score_lead(signals: list[str], has_phone: bool, has_address: bool) -> int:
    """A transparent 0-100 sum of signal weights plus contact completeness.

    Deliberately simple and explainable: this orders a review queue, it does not
    predict anything.
    """
    total = sum(SIGNALS.get(key, ('', 0))[1] for key in signals)
    total += 20 if has_phone else 0
    total += 10 if has_address else 0
    return max(0, min(100, total))


def parse_csv(content: bytes, overrides: dict[str, str] | None = None):
    """Parse an uploaded CSV into normalized lead dicts, warnings and column meta."""
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = content.decode('latin-1')

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return [], ['The file has no header row.'], _empty_meta()
    return _read_rows(list(reader.fieldnames), reader, overrides)


def _text(cell: object) -> str:
    """One spreadsheet cell as text.

    Whole numbers come back as floats often enough that 3125550142.0 would
    otherwise be read as an eleven-digit phone number.
    """
    if cell is None:
        return ''
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def parse_workbook(content: bytes, overrides: dict[str, str] | None = None,
                   sheet_name: str = ''):
    """Parse one lead-like sheet of an .xlsx using the same column rules.

    CRM workbooks often begin with a Read Me or summary sheet. Select the first
    sheet whose header contains a recognised phone or property-address column,
    but never merge multiple sheets: a lead list remains one explicit table.
    If no sheet is recognisable, retain the old first-nonempty-sheet behaviour
    so the normal mapping error remains useful.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - openpyxl is a pinned dependency
        return [], ['Excel files need the openpyxl package on the server.'], _empty_meta()

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - openpyxl raises many shapes
        return [], [f'The workbook could not be opened: {error}'], _empty_meta()

    try:
        worksheet_names = list(workbook.sheetnames)
        selected = None
        fallback = None
        # A short scan also handles exports that put a title above the table.
        if sheet_name and sheet_name not in worksheet_names:
            meta = _empty_meta()
            meta['worksheets'] = worksheet_names
            return [], [f'The worksheet “{sheet_name}” was not found.'], meta

        candidates = [workbook[sheet_name]] if sheet_name else workbook.worksheets
        for sheet in candidates:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [_text(cell) for cell in row]
                if not any(values):
                    if row_number >= 50:
                        break
                    continue
                if fallback is None:
                    fallback = (sheet, row_number, values)
                mapped, _flags = _map_headers(values, overrides)
                if 'phone' in mapped or 'property_address' in mapped:
                    selected = (sheet, row_number, values)
                    break
                if row_number >= 50:
                    break
            if selected is not None:
                break

        selected = selected or fallback
        if selected is None:
            return [], ['The workbook has no header row.'], _empty_meta()

        sheet, header_row, header = selected

        records = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            values = [_text(cell) for cell in row]
            if not any(values):
                continue
            # Pad or trim so ragged rows still line up with the header.
            values += [''] * (len(header) - len(values))
            records.append(dict(zip(header, values[:len(header)])))
    finally:
        workbook.close()

    parsed = _read_rows(header, records, overrides)
    parsed[2]['worksheets'] = worksheet_names
    parsed[2]['selected_sheet'] = sheet.title
    return parsed


def parse_upload(content: bytes, filename: str, overrides: dict[str, str] | None = None,
                 sheet_name: str = ''):
    """Dispatch on file extension. Anything else is rejected by the route."""
    if filename.lower().endswith('.xlsx'):
        return parse_workbook(content, overrides, sheet_name)
    return parse_csv(content, overrides)


def _empty_meta() -> dict:
    return {
        'columns': [], 'mapping': {}, 'signal_columns': {},
        'worksheets': [], 'selected_sheet': None,
    }


def _read_rows(fieldnames: list[str], rows: Iterable[dict],
               overrides: dict[str, str] | None = None):
    """Map header-aligned dict rows onto normalized leads. Shared by CSV and XLSX."""
    is_lofty = any(_key(column) == 'loftyleadid' for column in fieldnames)
    effective_overrides = dict(overrides or {})
    # "Lead Type" precedes "Tags" in Lofty exports, but it describes the CRM
    # contact type (Seller/Buyer), not a property signal. Tags can contain real
    # supported signals such as Expired and FSBO.
    if is_lofty and 'signals' not in effective_overrides and 'Tags' in fieldnames:
        effective_overrides['signals'] = 'Tags'
    fields, flags = _map_headers(list(fieldnames), effective_overrides)
    meta = {
        'columns': list(fieldnames),
        'mapping': {field: fields.get(field) for field in MAPPABLE_FIELDS},
        'signal_columns': dict(flags),
    }
    if 'phone' not in fields and 'property_address' not in fields:
        return [], ['No phone or property address column was recognised. '
                    'Choose which columns hold the phone number and the property address.'], meta

    leads: list[dict] = []
    warnings: list[str] = []
    for number, row in enumerate(rows, start=2):
        def cell(field: str) -> str:
            column = fields.get(field)
            return str(row.get(column) or '').strip() if column else ''

        name = cell('owner_name')
        if not name:
            name = ' '.join(part for part in (cell('first_name'), cell('last_name')) if part).strip()

        raw_phones = [cell('phone'), cell('phone_2'), cell('phone_3')]
        phones: list[str] = []
        for raw_phone in raw_phones:
            normalized = normalize_phone(raw_phone)
            if normalized and normalized not in phones:
                phones.append(normalized)
        phone = phones[0] if phones else None
        address = cell('property_address')
        # When a vendor splits state/ZIP out, build one geocodable address.
        # A city-only companion column remains `area` and does not unexpectedly
        # alter older files that already carry a complete address.
        if address and (cell('state') or cell('postal_code')):
            parts = [address]
            current = address.lower()
            for part in (cell('area'), cell('state'), cell('postal_code')):
                if part and part.lower() not in current:
                    parts.append(part)
            address = ', '.join(parts)
        if not phone and not address:
            if any(str(value or '').strip() for value in row.values()):
                warnings.append(f'Row {number}: no usable phone or property address — skipped.')
            continue
        if cell('phone') and not normalize_phone(cell('phone')):
            warnings.append(f'Row {number}: "{cell("phone")}" is not a usable phone number.')

        signals = parse_signals(cell('signals'))
        for signal, column in flags.items():
            if str(row.get(column) or '').strip().lower() in _TRUTHY and signal not in signals:
                signals.append(signal)

        details = parse_details(row.get(fields['details'])) if 'details' in fields else {}
        dnc = bool(cell('ok_to_text') and cell('ok_to_text').strip().lower() not in _TRUTHY)
        if is_lofty:
            for column, key in _LOFTY_DETAIL_COLUMNS.items():
                value = str(row.get(column) or '').strip()
                if value:
                    details[key] = value
        # Contact columns use the same existing details contract for every
        # vendor, not only known Lofty exports.
        emails = [cell('email_1'), cell('email_2')]
        emails = list(dict.fromkeys(value for value in emails if value))
        if emails:
            details['emails'] = emails
        if phones:
            details['_phone_numbers'] = [
                {'phone': value, 'dnc': dnc, 'owner_name': name}
                for value in phones
            ]

        # Preserve all other source fields exactly as labelled by the vendor.
        # Core/mapped columns stay authoritative and are not duplicated here.
        used_columns = set(fields.values()) | set(flags.values())
        if is_lofty:
            used_columns.update(_LOFTY_DETAIL_COLUMNS)
        imported_fields = {
            str(column): value
            for column in fieldnames
            if column not in used_columns
            if (value := _text(row.get(column)))
        }
        if imported_fields:
            details['_imported_fields'] = imported_fields
        for signal in signals_from_details(details):
            if signal not in signals:
                signals.append(signal)

        leads.append({
            'owner_name': name[:255] or None,
            'phone': phone,
            'property_address': address[:500] or None,
            'area': cell('area')[:120] or None,
            'signals': signals,
            'outreach_reason': cell('outreach_reason')[:300] or None,
            'details': details,
            'dnc': dnc,
            'score': score_lead(signals, bool(phone), bool(address)),
        })

    if not leads and not warnings:
        warnings.append('The file contained no data rows.')
    return leads, warnings, meta
